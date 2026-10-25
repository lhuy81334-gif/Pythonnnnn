import sys
import json
import time
import datetime
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import config
from features import extract_handcrafted_features
from ip_extractor import get_ip_location

if sys.platform.startswith('win'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def export_logs_to_excel(output_file: str = "Bao_Cao_Phat_Hien_Spam.xlsx") -> str:
    """
    Xuất File Excel chuẩn gồm 4 Sheet rõ ràng:
    - Sheet 1 (1_Thong_So_Dac_Trung): Bảng các thông số đo lường học máy.
    - Sheet 2 (2_Ket_Qua_Random_Forest): Bảng kết quả AI + Khung giải thích nguyên lý toán học.
    - Sheet 3 (3_Danh_Sach_Chi_Tiet): Nhật ký an ninh chi tiết.
    - Sheet 4 (4_Bang_Doi_Chieu_Email): BẢNG ĐỐI CHIẾU (STT <-> Tên email thật <-> Người gửi <-> Kịch bản thực tế).
    """
    if not config.DATABASE_LOG_PATH.exists():
        print("[!] Không tìm thấy dữ liệu nhật ký để xuất Excel.")
        return ""

    try:
        with open(config.DATABASE_LOG_PATH, "r", encoding="utf-8") as f:
            logs = json.load(f)
    except Exception as e:
        print(f"[!] Lỗi đọc file log: {e}")
        return ""

    if not logs:
        print("[!] Nhật ký email đang trống.")
        return ""

    raw_features = []
    rf_results = []
    detailed_records = []
    mapping_records = []

    for idx, item in enumerate(logs, 1):
        dt_str = datetime.datetime.fromtimestamp(item.get("timestamp", 0)).strftime("%H:%M:%S %d/%m/%Y")
        subject = item.get("subject", "")
        body = item.get("body", "")
        sender = item.get("sender", "")
        attachments = item.get("attachments", [])
        indicators = item.get("indicators", {})
        if not indicators:
            indicators = extract_handcrafted_features(subject, body, attachments)
            
        has_crypto_or_money = int(bool(indicators.get("has_crypto", 0) or indicators.get("has_huge_money", 0)))
        has_malware = int(bool(indicators.get("has_dangerous_file", 0) or indicators.get("has_double_extension", 0)))
        is_flood = int(bool(item.get("is_flood_attack", False)))
        spam_prob = round(item.get("spam_score", 0) * 100, 2)
        label = 1 if item.get("is_spam") else 0

        # Xác định tên kịch bản thực tế
        if "legacy" in subject.lower() or "cancer" in body.lower():
            scenario_name = "Kịch bản 1: Lừa đảo Crypto / Thừa kế (User/Pass + USDT)"
        elif "100k" in subject.lower() or "nhà cái" in subject.lower():
            scenario_name = "Kịch bản 2: Spam Quảng cáo Khuyến mãi / Nhà cái"
        elif "họp nhóm" in subject.lower() or "đồ án" in subject.lower():
            scenario_name = "Kịch bản 3: Email Học tập / Công việc Hợp lệ (Sạch)"
        elif "500 usdt" in subject.lower():
            scenario_name = "Kịch bản 4: Tấn công dồn dập Flood / Sybil Attack (Bắn dồn 4 đợt)"
        elif has_malware or "lương" in subject.lower():
            scenario_name = "Kịch bản 5: Email chứa File độc hại (.pdf.exe)"
        else:
            scenario_name = "Email Thử Nghiệm Kết Nối"

        # Sheet 1: Thông số đo lường
        raw_features.append({
            "STT": idx,
            "Số Lượng Link URL": indicators.get("url_count", 0),
            "Bẫy Tài Khoản (User/Pass)": indicators.get("has_credentials", 0),
            "Dấu Hiệu Tiền Ảo (Crypto/USDT)": has_crypto_or_money,
            "File Đính Kèm Độc Hại (.exe)": has_malware,
            "Tấn Công Dồn Dập (Flood)": is_flood
        })

        # Sheet 2: Kết quả Random Forest
        rf_results.append({
            "STT": idx,
            "Số Lượng Link URL": indicators.get("url_count", 0),
            "Bẫy Tài Khoản": indicators.get("has_credentials", 0),
            "Tiền Ảo / USDT": has_crypto_or_money,
            "File Độc Hại": has_malware,
            "Tấn Công Flood": is_flood,
            "Điểm Rủi Ro Random Forest (%)": f"{spam_prob}%",
            "Nhãn Kết Quả (0=Sạch, 1=Spam)": label
        })

        sender_ip_val = item.get("sender_ip", "Không xác định")
        sender_location = get_ip_location(sender_ip_val)

        # Sheet 3: Danh sách thông tin chi tiết
        detailed_records.append({
            "STT": idx,
            "Thời Gian Nhận": dt_str,
            "Người Gửi": sender,
            "IP Nguồn (Sender IP)": sender_ip_val,
            "Vị Trí / Quốc Gia": sender_location,
            "Tiêu Đề": subject,
            "Nội Dung Tóm Tắt": body[:120] + ("..." if len(body) > 120 else ""),
            "File Đính Kèm": ", ".join(attachments) if attachments else "Không có",
            "Phân Loại AI": item.get("verdict", "SPAM" if item.get("is_spam") else "HỢP LỆ"),
            "Điểm Rủi Ro (%)": f"{spam_prob}%",
            "Cảnh Báo Flood Attack": "CÓ (BẤT THƯỜNG)" if is_flood else "KHÔNG",
            "Lý Do Cảnh Báo": "; ".join(item.get("flood_reasons", [])) if item.get("flood_reasons") else "Không có"
        })

        # Sheet 4: Bảng Đối Chiếu (Mapping rõ ràng từng dòng là email gì)
        mapping_records.append({
            "STT trong Dataset": idx,
            "Tên Email Thực Tế": subject,
            "Tài Khoản Gửi Thật": sender,
            "IP & Quốc Gia": f"{sender_ip_val} ({sender_location})",
            "Kịch Bản Thực Chiến": scenario_name,
            "Đặc Trưng Bị Nhận Diện": f"URL={indicators.get('url_count',0)}, UserPass={indicators.get('has_credentials',0)}, Crypto={has_crypto_or_money}, Flood={is_flood}",
            "Kết Luận AI": "🚨 SPAM / PHISHING" if label == 1 else "✅ EMAIL SẠCH"
        })

    df_sheet1 = pd.DataFrame(raw_features)
    df_sheet2 = pd.DataFrame(rf_results)
    df_sheet3 = pd.DataFrame(detailed_records)
    df_sheet4 = pd.DataFrame(mapping_records)

    # Thử lưu file chính hoặc các file thay thế nếu đang mở trong Excel
    for attempt_name in [output_file, "Bao_Cao_Phat_Hien_Spam_Final.xlsx", "Bao_Cao_Phat_Hien_Spam_Moi.xlsx"]:
        try:
            target_path = Path(attempt_name)
            with pd.ExcelWriter(target_path, engine='openpyxl') as writer:
                df_sheet1.to_excel(writer, sheet_name="1_Thong_So_Dac_Trung", index=False)
                df_sheet2.to_excel(writer, sheet_name="2_Ket_Qua_Random_Forest", index=False)
                df_sheet3.to_excel(writer, sheet_name="3_Danh_Sach_Chi_Tiet", index=False)
                df_sheet4.to_excel(writer, sheet_name="4_Bang_Doi_Chieu_Email", index=False)
            target_file = target_path
            break
        except PermissionError:
            continue
    else:
        # Nếu tất cả đều bị mở, tạo theo timestamp
        target_file = Path(f"Bao_Cao_Phat_Hien_Spam_{int(time.time())}.xlsx")
        with pd.ExcelWriter(target_file, engine='openpyxl') as writer:
            df_sheet1.to_excel(writer, sheet_name="1_Thong_So_Dac_Trung", index=False)
            df_sheet2.to_excel(writer, sheet_name="2_Ket_Qua_Random_Forest", index=False)
            df_sheet3.to_excel(writer, sheet_name="3_Danh_Sach_Chi_Tiet", index=False)
            df_sheet4.to_excel(writer, sheet_name="4_Bang_Doi_Chieu_Email", index=False)


    # Định dạng khung giải thích Random Forest vào Sheet 2
    wb = openpyxl.load_workbook(target_file)
    ws2 = wb["2_Ket_Qua_Random_Forest"]
    
    start_row = len(rf_results) + 3
    ws2.cell(row=start_row, column=1, value="📌 GIẢI THÍCH NGUYÊN LÝ RANDOM FOREST RA KẾT QUẢ:").font = Font(bold=True, size=11, color="1F4E79")
    
    explanation_text = (
        "1. Cơ chế hoạt động: Rừng ngẫu nhiên gồm 150 Cây Quyết Định (Decision Trees) làm việc độc lập.\n"
        "2. Biểu quyết (Majority Voting): Mỗi cây kiểm tra các đặc trưng (Link, User/Pass, Tiền ảo USDT, File .exe, Tần suất Flood) và bỏ 1 phiếu (0 hoặc 1).\n"
        "3. Công thức tính điểm: Điểm Rủi Ro (%) = (Tổng số cây bỏ phiếu Spam / 150) * 100%.\n"
        "4. Ra quyết định: Nếu Điểm Rủi Ro >= 50% -> Gắn Nhãn 1 (Spam / Bất Thường / Mã Độc); Nếu < 50% -> Gắn Nhãn 0 (Hợp Lệ / Sạch)."
    )
    
    ws2.merge_cells(start_row=start_row+1, start_column=1, end_row=start_row+5, end_column=8)
    cell_explain = ws2.cell(row=start_row+1, column=1, value=explanation_text)
    cell_explain.alignment = Alignment(wrap_text=True, vertical="top")
    cell_explain.fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    cell_explain.font = Font(size=10, italic=False)

    wb.save(target_file)
    print(f"[+] ĐÃ XUẤT THÀNH CÔNG 4 SHEET ĐẦY ĐỦ VÀO: {target_file.resolve()}")
    return str(target_file.resolve())

if __name__ == "__main__":
    export_logs_to_excel()
